import streamlit as st
import pandas as pd
import openpyxl
from io import BytesIO
import pdfplumber
import re

# ==========================================
# CONFIGURAÇÃO DA PÁGINA
# ==========================================
st.set_page_config(page_title="Gerador Financeiro UFSM", page_icon="💰", layout="wide")

# ==========================================
# LISTAS E DADOS FIXOS
# ==========================================
LISTA_DIARIAS = ["Diárias no país", "Diárias no exterior", "Auxílio para desenvolvimento de estudos e pesquisas", "Diárias a colaboradores eventuais no país"]
LISTA_PF = ["Direitos autorais", "Serviços técnicos profissionais", "Serviços de limpeza e conservação", "Serviços de apoio administrativo, técnico e operacional", "Obrigações Tributárias e Contributivas (cota patronal 20%)"]
LISTA_PASSAGENS = ["Passagens para o país", "Passagens para o exterior", "Locação de meios de transportes", "Locomoção urbana", "Outras despesas com locomoção"]
LISTA_OBRAS = ["Estudos e Projetos", "Obras em andamento", "Instalações", "Almoxarifado de obras", "Outras obras e instalações"]
LISTA_PJ = [
    "Assinaturas de periódicos e anuidades", "Direitos autorais", "Serviços técnicos profissionais", "Manutenção de software", "Locação de imóveis", "Locação de softwares", "Locação de máquinas e equipamentos", "Locação de bens Mov. Out. naturezas e intangíveis", "Manutenção e conservação de bens imóveis", "Manutenção e conservação de máquinas e equipamentos", "Serviço de estacionamento de veículos", "Manutenção e conservação de veículos", "Exposições, congressos e conferências", "Confecção de uniformes", "Desenvolvimento de software", "Suporte de infraestrutura de TI", "Suporte a usuários de TI", "Hospedagem de sistemas", "Locação de equipamentos de processamento de dados", "Fornecimento de alimentação", "Serviços de energia elétrica", "Serviços de água e esgoto", "Serviços de comunicação em geral", "Serviços médico-hospitalar, odontológicos e laboratoriais", "Serviços de análises e pesquisas científicas", "Serviços de tecnologia da informação", "Serviços de telecomunicações", "Serviços de áudio, vídeo e foto", "Serviços de produção industrial", "Serviços gráficos e editoriais", "Seguros em geral", "Confecção de material de acondicionamento e embalagem", "Vale-transporte", "Fretes e transportes de encomendas", "Serviço de apoio administrativo, técnico e operacional", "Hospedagens", "Serviços de cópias e reprodução de documentos", "Serviços de publicidade legal", "Aquisição de softwares sob encomenda", "Manutenção e conservação de equip. de processamento de dados", "Comunicação de dados", "Testes e confecções de placas relacionadas ao objeto do projeto", "Outros serviços de terceiros pessoa jurídica"
]
LISTA_CONSUMO = [
    "Combustíveis e lubrificantes automotivos", "Gás e outros materiais engarrafados", "Alimentos para animais", "Gêneros de alimentação", "Animais para pesquisa e abate", "Material farmacológico", "Material odontológico", "Material químico", "Material educativo e esportivo", "Material de expediente", "Material de processamento de dados", "Materiais e medicamentos para uso veterinário", "Material de acondicionamento e embalagem", "Material de copa e cozinha", "Material de limpeza e produtos de higienização", "Uniformes, tecidos e aviamentos", "Material para manutenção de bens imóveis/instalações", "Material para manutenção de bens móveis", "Material elétrico e eletrônico", "Material de proteção e segurança", "Material para áudio, vídeo e foto", "Material para comunicações", "Sementes, mudas de plantas e insumos", "Material para produção industrial", "Material laboratorial", "Material hospitalar", "Material para manutenção de veículos", "Material biológico", "Material para utilização em gráfica", "Ferramentas", "Material de sinalização visual e outros", "Material bibliográfico", "Aquisição de software - produto", "Material para divulgação", "Materiais de Consumo para utilização no laboratório", "Outros Materiais de Consumo"
]

FONTES_OPCOES = [
    "de investimento em pesquisa, previsto no projeto de prestação de serviços abaixo:",
    "do repasse de terceiros, interessados em resultados da pesquisa",
    "do pagamento de clientes, pelos serviços oferecidos pelo projeto",
    "da venda de subprodutos decorrentes de atividades práticas acadêmicas da UFSM",
    "da venda de produtos produzidos através do próprio projeto",
    "de patrocínios e taxas de inscrição em evento(s)",
    "de investimento de recursos próprios, através do orçamento da UFSM",
    "do compartilhamento de infraestrutura",
    "do repasse de outras fontes de recursos orçamentários da União (TED)"
]

if 'df_equipe' not in st.session_state: 
    st.session_state.df_equipe = pd.DataFrame(columns=[
        "Vinculado à UFSM?", "Tipo Remuneração", "Nome", "SIAPE ou Forma Contratação", 
        "CPF", "Carga Horária", "Nº Pagamentos", "Valor Parcela (R$)"
    ])
if 'equip' not in st.session_state: st.session_state.equip = []

col_config_dinheiro = {
    "Valor Parcela": st.column_config.NumberColumn(format="R$ %.2f"),
    "Total": st.column_config.NumberColumn(format="R$ %.2f"),
    "Valor Unitário": st.column_config.NumberColumn(format="R$ %.2f")
}

total_base_infra = 0.0 
total_obras_equip = 0.0

fundacao_escolhida = st.sidebar.selectbox(
    "🏛️ Selecione a Fundação",
    ["FATEC", "FAURGS", "FUNDEP", "FDMS"],
    help="As taxas e cálculos são ajustados automaticamente de acordo com as regras de cada fundação."
)
st.sidebar.divider()

st.title("💰 Gerador de Dados Financeiros (Plano de Trabalho)")

st.markdown("### 📄 Passo 0: Autopreenchimento da Equipe (Opcional)")
arquivo_pdf = st.file_uploader("Insira o Relatório do Projeto (.PDF) para extrair os nomes e SIAPEs da equipe automaticamente:", type=["pdf"])

if arquivo_pdf:
    with st.spinner("Lendo participantes do PDF..."):
        try:
            texto_limpo = ""
            with pdfplumber.open(arquivo_pdf) as pdf:
                for page in pdf.pages:
                    ext = page.extract_text()
                    if ext: texto_limpo += ext + "\n"
            
            m_inicio = re.search(r'PARTICIPANTES', texto_limpo, re.IGNORECASE)
            if m_inicio:
                idx = m_inicio.end()
                end_idx = len(texto_limpo)
                for f in [r'UNIDADES VINCULADAS\s*\n', r'CLASSIFICAÇÕES', r'REGIÕES DE ATUAÇÃO']:
                    mf = re.search(f, texto_limpo[idx:], re.IGNORECASE)
                    if mf:
                        pos = idx + mf.start()
                        if pos < end_idx: end_idx = pos
                bloco = texto_limpo[idx:end_idx].strip()
                
                matches = re.finditer(r'(\d{5,15})\s*-\s*([A-ZÀ-Ÿ\s\']+?)\s*(?=[A-ZÀ-Ÿ][a-zà-ÿ]|UNIDADES VINCULADAS|CLASSIFICAÇÕES|$)', bloco)
                
                participantes_extraidos = []
                for match in matches:
                    siape_pdf = match.group(1).strip()
                    nome_limpo = re.sub(r'\s+', ' ', match.group(2).strip()).strip()
                    corte_idx = len(nome_limpo)
                    
                    for p in ["VÍNCULO", "VINCULO", "CURSO", "LOTAÇÃO", "LOTACAO", "FUNÇÃO", "FUNCAO"]:
                        idx_p = nome_limpo.upper().find(p)
                        if idx_p != -1 and idx_p < corte_idx: corte_idx = idx_p
                        
                    nome_final = nome_limpo[:corte_idx].strip(" -/")
                    if len(nome_final) > 2: 
                        participantes_extraidos.append({"nome": nome_final, "siape": siape_pdf})
                
                participantes_unicos = []
                nomes_vistos = set()
                for p in participantes_extraidos:
                    if p["nome"] not in nomes_vistos:
                        nomes_vistos.add(p["nome"])
                        participantes_unicos.append(p)
                
                if participantes_unicos:
                    df_atual = st.session_state.df_equipe
                    nomes_existentes = df_atual["Nome"].tolist() if not df_atual.empty else []
                    
                    novos_registros = []
                    for p in participantes_unicos:
                        if p["nome"] not in nomes_existentes:
                            novos_registros.append({
                                "Vinculado à UFSM?": True, "Tipo Remuneração": "", "Nome": p["nome"],
                                "SIAPE ou Forma Contratação": p["siape"], "CPF": "", "Carga Horária": 0,
                                "Nº Pagamentos": 1, "Valor Parcela (R$)": 0.0
                            })
                    
                    if novos_registros:
                        df_novos = pd.DataFrame(novos_registros)
                        st.session_state.df_equipe = pd.concat([df_atual, df_novos], ignore_index=True)
                        st.success(f"✅ {len(novos_registros)} participantes (com SIAPE) encontrados e adicionados à tabela abaixo!")
            else:
                st.warning("Nenhum bloco de participantes encontrado no PDF.")
        except Exception as e:
            st.error(f"Erro na leitura do PDF: {e}")

st.divider()

st.header("1. Equipe Executora")
st.write("Preencha os dados abaixo. Desmarque a caixinha **'Vinculado à UFSM?'** para colaboradores externos. Você pode adicionar novas pessoas clicando na última linha em branco.")

configuracao_colunas = {
    "Vinculado à UFSM?": st.column_config.CheckboxColumn("Vinculado à UFSM?", default=True, width="small"),
    "Valor Parcela (R$)": st.column_config.NumberColumn("Valor Parcela (R$)", format="R$ %.2f", min_value=0.0),
    "Carga Horária": st.column_config.NumberColumn("Carga Horária", min_value=0, step=1),
    "Nº Pagamentos": st.column_config.NumberColumn("Nº Pagamentos", min_value=1, step=1)
}

df_editado = st.data_editor(
    st.session_state.df_equipe,
    num_rows="dynamic",
    use_container_width=True,
    hide_index=True,
    column_config=configuracao_colunas,
    key="editor_equipe"
)

st.session_state.df_equipe = df_editado
st.session_state.eq_vinc = []
st.session_state.eq_nao_vinc = []

for index, row in df_editado.iterrows():
    nome = str(row.get("Nome", "")).strip()
    if not nome or nome == "nan": continue
    
    try: pagamentos = int(row.get("Nº Pagamentos", 1))
    except: pagamentos = 1
    try: valor = float(row.get("Valor Parcela (R$)", 0.0))
    except: valor = 0.0
    try: ch = int(row.get("Carga Horária", 0))
    except: ch = 0
    
    total_participante = pagamentos * valor
    
    participante = {
        "Tipo Remuneração": str(row.get("Tipo Remuneração", "")).strip(),
        "Nome": nome,
        "SIAPE/MAT" if row.get("Vinculado à UFSM?", True) else "Forma Contratação": str(row.get("SIAPE ou Forma Contratação", "")).strip(),
        "CPF": str(row.get("CPF", "")).strip(),
        "Carga Horária": ch,
        "Nº Pagamentos": pagamentos,
        "Valor Parcela": valor,
        "Total": total_participante
    }
    
    if row.get("Vinculado à UFSM?", True):
        st.session_state.eq_vinc.append(participante)
    else:
        st.session_state.eq_nao_vinc.append(participante)

total_base_infra += sum(item["Total"] for item in st.session_state.eq_vinc)
total_base_infra += sum(item["Total"] for item in st.session_state.eq_nao_vinc)

st.header("2. Despesas e Serviços")
def renderizar_tabela_fixa(titulo, lista_itens, prefixo_chave):
    valores = {}
    with st.expander(titulo):
        for item in lista_itens:
            val = st.number_input(f"{item} (R$)", min_value=0.0, step=50.0, key=f"{prefixo_chave}_{item}")
            if val > 0: valores[item] = val
    return valores

dados_diarias = renderizar_tabela_fixa("4.2 - Diárias", LISTA_DIARIAS, "diaria")
dados_pj = renderizar_tabela_fixa("4.3 - Serviços de Terceiros (PJ)", LISTA_PJ, "pj")
dados_pf = renderizar_tabela_fixa("4.4 - Serviços de Terceiros (PF)", LISTA_PF, "pf")
dados_passagens = renderizar_tabela_fixa("4.5 - Passagens e Locomoção", LISTA_PASSAGENS, "pass")
dados_consumo = renderizar_tabela_fixa("4.6 - Material de Consumo", LISTA_CONSUMO, "cons")
dados_obras = renderizar_tabela_fixa("4.8 - Obras e Instalações (Isento de Taxa UFSM)", LISTA_OBRAS, "obras")

for d in [dados_diarias, dados_pj, dados_pf, dados_passagens, dados_consumo]: 
    total_base_infra += sum(d.values())
total_obras_equip += sum(dados_obras.values())

st.header("3. Anexo I - Material Permanente")
with st.expander("Equipamento Permanente (Isento de Taxa UFSM)", expanded=True):
    with st.form("form_equip", clear_on_submit=True):
        c1, c2, c3 = st.columns([3, 1, 1])
        desc_eq = c1.text_input("Especificação")
        qtd_eq = c2.number_input("Quantidade", min_value=1, step=1)
        val_eq = c3.number_input("Valor Unitário (R$)", min_value=0.0, step=100.0)
        
        if st.form_submit_button("Adicionar Equipamento") and desc_eq:
            st.session_state.equip.append({"Especificação": desc_eq, "Quantidade": qtd_eq, "Valor Unitário": val_eq, "Total": qtd_eq * val_eq})
            st.rerun()
            
    if st.session_state.equip: 
        st.dataframe(st.session_state.equip, column_config=col_config_dinheiro, hide_index=True, use_container_width=True)

total_obras_equip += sum(item["Total"] for item in st.session_state.equip)
st.divider()

st.header("4. Fontes e Usos (Seção 3)")
with st.expander("3.1 - FONTES (Especificação dos Recursos)", expanded=True):
    st.write("Marque com um 'X' (selecione) as origens do dinheiro deste projeto. O valor total do projeto será associado a elas no formulário final.")
    escolhas_fontes = {}
    
    for op in FONTES_OPCOES:
        chk = st.checkbox(op)
        escolhas_fontes[op] = chk
        
        if chk and "previsto no projeto de prestação de serviços abaixo" in op:
            st.info("Informe os dados do projeto de prestação de serviços que aportará recursos:")
            titulo_fonte = st.text_input("Título do Projeto de Prestação de Serviços (Fonte):")
            registro_fonte = st.text_input("Nº de Registro (Fonte):")
        else:
            if "previsto no projeto de prestação de serviços abaixo" in op:
                titulo_fonte = ""
                registro_fonte = ""

with st.expander("4 - PLANO DE APLICAÇÃO (Apenas Prestação de Serviços)", expanded=True):
    st.write("Marque a caixa abaixo SOMENTE se este projeto atual que você está cadastrando vai investir/aportar recursos financeiros em OUTRO projeto.")
    aplica_recursos = st.checkbox("✅ Este projeto atual investirá/aportará recursos em outro projeto?")
    if aplica_recursos:
        st.info("Informe os dados e o VALOR do projeto (Pesquisa, Ensino ou Extensão) que RECEBERÁ os recursos:")
        col_a1, col_a2 = st.columns(2)
        titulo_aplicacao = col_a1.text_input("Título do projeto recebedor:")
        registro_aplicacao = col_a2.text_input("Nº de Registro do projeto recebedor:")
        valor_aplicacao = st.number_input("Valor a ser investido/repassado (R$):", min_value=0.0, step=100.0)
    else:
        titulo_aplicacao = ""
        registro_aplicacao = ""
        valor_aplicacao = 0.0

with st.expander("6 - CRONOGRAMA DE DESEMBOLSO", expanded=True):
    st.write("Escolha apenas UMA opção de período e preencha os valores.")
    tipo_cronograma = st.radio("Período de Desembolso:", ["Mensal", "Semestral", "Anual"], horizontal=True)
    qtd_periodos = st.number_input(f"Quantidade de períodos ({tipo_cronograma.lower()}s):", min_value=1, max_value=60, value=1)
    
    st.write("### Preencha os valores de cada período:")
    cronograma_dados = []
    colunas_cronograma = st.columns(3)
    for i in range(int(qtd_periodos)):
        val_periodo = colunas_cronograma[i % 3].number_input(f"{tipo_cronograma} {i+1} (R$)", min_value=0.0, step=100.0, key=f"crono_{i}")
        cronograma_dados.append({"Período": f"{tipo_cronograma} {i+1}", "Valor": val_periodo})


# ==========================================
# CÁLCULOS MATEMÁTICOS FINAIS
# ==========================================
subtotal_projeto = total_base_infra + total_obras_equip

taxa_ufsm = 0.08 if total_base_infra > 200000 else 0.05
valor_infra_ufsm = total_base_infra * taxa_ufsm

if fundacao_escolhida in ["FATEC", "FDMS"]:
    valor_taxa_fundacao = subtotal_projeto * 0.10
else:
    valor_taxa_fundacao = ((subtotal_projeto + valor_infra_ufsm) / 0.9) * 0.10

total_geral_final = subtotal_projeto + valor_infra_ufsm + valor_taxa_fundacao

# ==========================================
# RENDERIZAÇÃO DA BARRA LATERAL (Cálculos de 3.2 - USOS)
# ==========================================
with st.sidebar:
    st.title("3.2 - USOS (Resumo)")
    
    st.caption("Despesas de Custeio (Equipe + Serviços)")
    st.write(f"R$ {total_base_infra:,.2f}")
    
    st.caption("Despesas de Capital (Obras e Equipamentos)")
    st.write(f"R$ {total_obras_equip:,.2f}")
    
    st.caption(f"Ressarcimento Infraestrutura UFSM ({int(taxa_ufsm*100)}%)", 
               help="Regra da UFSM: O ressarcimento é de 5% para custeios de até 200 mil reais, e passa a ser de 8% para custeios acima de 200 mil reais. Obras e Equipamentos (Anexo I) são totalmente isentos desta taxa.")
    st.write(f"R$ {valor_infra_ufsm:,.2f}")
    
    tipo_calculo = "Cálculo Direto" if fundacao_escolhida in ["FATEC", "FDMS"] else "Gross-up"
    st.caption(f"Despesas Operacionais {fundacao_escolhida} (10% - {tipo_calculo})")
    st.write(f"R$ {valor_taxa_fundacao:,.2f}")
    
    st.divider()
    st.metric("TOTAL GERAL DO PROJETO", f"R$ {total_geral_final:,.2f}")
    st.divider()
    
    if st.button("📥 GERAR ARQUIVO", type="primary", use_container_width=True):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        def criar_aba_dinamica(titulo, dados):
            ws = wb.create_sheet(titulo)
            if dados:
                ws.append(list(dados[0].keys()))
                for linha in dados: ws.append(list(linha.values()))
            else:
                ws.append(["Nenhum item cadastrado"])

        def criar_aba_fixa(titulo, dicionario):
            ws = wb.create_sheet(titulo)
            ws.append(["Categoria", "Valor (R$)"])
            if dicionario:
                for cat, val in dicionario.items(): ws.append([cat, val])
            else:
                ws.append(["Nenhum item preenchido"])

        ws_config = wb.create_sheet("Config_Raichu")
        ws_config.append(["Fundacao_Escolhida", fundacao_escolhida])
        ws_config.append(["Cronograma_Tipo", tipo_cronograma])
        ws_config.append(["Total_Geral", total_geral_final])
        
        criar_aba_dinamica("Equipe_Vinc", st.session_state.eq_vinc)
        criar_aba_dinamica("Equipe_Nao_Vinc", st.session_state.eq_nao_vinc)
        criar_aba_fixa("Diarias", dados_diarias)
        criar_aba_fixa("Servicos_PJ", dados_pj)
        criar_aba_fixa("Servicos_PF", dados_pf)
        criar_aba_fixa("Passagens", dados_passagens)
        criar_aba_fixa("Consumo", dados_consumo)
        criar_aba_fixa("Obras", dados_obras)
        criar_aba_dinamica("Anexo_1", st.session_state.equip)

        ws_fontes = wb.create_sheet("Fontes_3.1")
        ws_fontes.append(["Fonte", "Selecionado", "Título (Se Prestação)", "Registro (Se Prestação)"])
        for op, checked in escolhas_fontes.items():
            if checked:
                if "previsto no projeto de prestação" in op:
                    ws_fontes.append([op, "X", titulo_fonte, registro_fonte])
                else:
                    ws_fontes.append([op, "X", "", ""])

        ws_aplicacao = wb.create_sheet("Aplicacao_4")
        ws_aplicacao.append(["Possui_Aporte", "Titulo_Recebedor", "Registro_Recebedor", "Valor_Aplicacao"])
        if aplica_recursos:
            ws_aplicacao.append(["Sim", titulo_aplicacao, registro_aplicacao, valor_aplicacao])
        else:
            ws_aplicacao.append(["Nao", "", "", 0.0])

        criar_aba_dinamica("Cronograma_6", cronograma_dados)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        st.success("Arquivo gerado! Envie-o para o preenchimento automático no Raichu.")
        st.download_button(
            label="📥 Baixar Dados_Financeiros.xlsx",
            data=output,
            file_name="Dados_Financeiros.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary"
        )
